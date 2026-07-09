from dataclasses import dataclass
from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.base import Base
from app.indexing.pipeline import IndexingPipeline
from app.indexing.sqlite_fts import SQLiteFtsIndex
from app.llm.provider import ModelInfo
from app.main import create_app
from app.models.chunk import Chunk
from app.models.document import Document
from app.repositories.sources import SourceRepository


ANCHORS = {
    "markdown": "ragmarkertopic",
    "text": "plainnotetopic",
    "html": "htmlworkflowtopic",
    "pdf": "pdfcitationtopic",
    "docx": "docxplanningtopic",
}


class FakeChatModelClient:
    def generate_answer(self, question, context) -> str:
        return f"Answer for {question} with {len(context.items)} source item(s)."


class FakeProvider:
    provider_id = "fake"

    def get_chat_client(self, model_id: str) -> FakeChatModelClient:
        return FakeChatModelClient()


@dataclass(frozen=True)
class FakeSelection:
    task: str
    provider: FakeProvider
    model: ModelInfo

    @property
    def full_name(self) -> str:
        return self.model.full_name


class FakeModelRouter:
    def __init__(self) -> None:
        self.provider = FakeProvider()
        self.model = ModelInfo(
            provider_id="fake",
            model_id="chat",
            display_name="Fake Chat",
            capabilities=["chat"],
        )

    def select_model(self, task: str) -> FakeSelection:
        return FakeSelection(task=task, provider=self.provider, model=self.model)


def test_realistic_folder_regression_indexes_multi_format_files_and_returns_traceable_citations(tmp_path) -> None:
    """Cover a local folder with multiple realistic file formats and fixed retrieval queries."""

    knowledge_dir = tmp_path / "Knowledge Root With Long Folder Name For Layout Regression"
    nested_dir = knowledge_dir / "nested" / "very-long-path-segment-for-regression"
    nested_dir.mkdir(parents=True)
    _write_regression_files(knowledge_dir, nested_dir)
    session_factory = _make_session_factory()
    session = session_factory()
    source = SourceRepository(session).create(
        source_type="local_directory",
        name="Phase 6 Regression Knowledge Source",
        uri=str(knowledge_dir),
        storage_mode="local_only",
        sync_direction="read_only",
    )
    lexical_index = SQLiteFtsIndex(session)

    job = IndexingPipeline(session, lexical_index=lexical_index).run_source_index(source.source_id)

    assert job.status == "completed"
    assert job.total_items == 5
    assert job.processed_items == 5
    assert job.failed_items == 0
    assert session.query(Document).count() == 5
    assert session.query(Chunk).count() >= 5
    formats = {
        document.metadata_json["source_format"]
        for document in session.query(Document).all()
    }
    assert formats == {"markdown", "text", "html", "pdf", "docx"}
    session.close()

    app = create_app()
    app.state.session_factory = session_factory
    app.state.model_router = FakeModelRouter()
    client = TestClient(app)

    for expected_format, query in ANCHORS.items():
        response = client.post("/search", json={"query": query, "top_k": 3})
        assert response.status_code == 200
        body = response.json()
        assert body["results"], f"expected a result for {query}"
        top_result = body["results"][0]
        assert top_result["document"]["metadata"]["source_format"] == expected_format
        assert top_result["citation"]["document_id"] == top_result["document_id"]
        assert top_result["citation"]["chunk_id"] == top_result["chunk_id"]

    chat_response = client.post("/chat", json={"message": "pdfcitationtopic", "top_k": 3})
    assert chat_response.status_code == 200
    chat_body = chat_response.json()
    assert chat_body["citations"]
    assert chat_body["citations"][0]["document_id"] > 0
    assert chat_body["citations"][0]["chunk_id"] > 0
    assert "pdfcitationtopic" in (chat_body["citations"][0]["snippet"] or "")


def _make_session_factory():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)


def _write_regression_files(root: Path, nested_dir: Path) -> None:
    (root / "01-rag-notes.md").write_text(
        "# RAG Notes\n\n"
        "The ragmarkertopic note explains grounded answers from local markdown notes.\n",
        encoding="utf-8",
    )
    (root / "02-plain-note.txt").write_text(
        "The plainnotetopic file records a lightweight text note for retrieval.\n",
        encoding="utf-8",
    )
    (nested_dir / "03-workflow-page.html").write_text(
        "<html><head><title>Workflow Page</title></head>"
        "<body><main><h1>Workflow</h1><p>The htmlworkflowtopic page has a long path.</p></main></body></html>",
        encoding="utf-8",
    )
    _write_pdf(
        nested_dir / "04-citation-source.pdf",
        ["The pdfcitationtopic page should be returned with a traceable citation."],
    )
    _write_docx(
        nested_dir / "05-planning-note.docx",
        "Planning Note",
        "The docxplanningtopic document validates Word ingestion in the regression folder.",
    )


def _write_pdf(path: Path, page_texts: list[str]) -> None:
    import fitz

    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def _write_docx(path: Path, heading: str, paragraph: str) -> None:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(paragraph)
    document.save(path)
